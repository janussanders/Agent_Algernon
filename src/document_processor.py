#!/usr/bin/env python3

import os
from typing import List, Optional
from pathlib import Path
from loguru import logger
from llama_index.core import (
    Settings,
    VectorStoreIndex,
    SimpleDirectoryReader,
    ServiceContext
)
from llama_index.vector_stores.qdrant import QdrantVectorStore
from llama_index.embeddings.huggingface import HuggingFaceEmbedding
from qdrant_client import QdrantClient, models
from qdrant_client.http.models import Distance, VectorParams
from llama_index.llms.sambanovasystems import SambaNovaCloud
from pypdf import PdfReader
import docx
import magic
from docx import Document
import openpyxl

class DocumentProcessor:
    """Process uploaded documents and extract text content"""
    
    def __init__(self, libmagic_path=None):
        # Initialize magic for MIME type detection
        self.mime = magic.Magic(mime=True)

    def extract_text(self, file_path: str) -> str:
        """Extract text from various document formats"""
        try:
            file_type = self.mime.from_file(file_path)
            
            if 'pdf' in file_type:
                return self._extract_pdf(file_path)
            elif 'msword' in file_type or 'officedocument' in file_type:
                return self._extract_docx(file_path)
            elif 'text' in file_type:
                return self._extract_txt(file_path)
            else:
                raise ValueError(f"Unsupported file type: {file_type}")
                
        except Exception as e:
            logger.error(f"Error extracting text: {str(e)}")
            return ""

    def _extract_pdf(self, file_path: str) -> str:
        """Extract text from a PDF file"""
        try:
            reader = PdfReader(file_path)
            text = ""
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text
            return text
        except Exception as e:
            logger.error(f"Error extracting text from PDF {file_path}: {str(e)}")
            return ""

    def _extract_docx(self, file_path: str) -> str:
        """Extract text from a Word document"""
        try:
            doc = Document(file_path)
            return "\n".join([para.text for para in doc.paragraphs])
        except Exception as e:
            logger.error(f"Error extracting text from Word document {file_path}: {str(e)}")
            return ""

    def _extract_txt(self, file_path: str) -> str:
        """Extract text from a plain text file"""
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                return file.read()
        except Exception as e:
            logger.error(f"Error extracting text from text file {file_path}: {str(e)}")
            return ""

    def _extract_excel(self, file_path: str) -> str:
        """Extract text from an Excel file"""
        try:
            workbook = openpyxl.load_workbook(file_path)
            text = ""
            for sheet in workbook:
                for row in sheet.iter_rows(values_only=True):
                    text += " ".join([str(cell) for cell in row if cell is not None]) + "\n"
            return text
        except Exception as e:
            logger.error(f"Error extracting text from Excel file {file_path}: {str(e)}")
            return ""

    def process_document(self, file_path):
        # Example method to process a document
        file_type = self.mime.from_file(file_path)
        print(f"The file type of {file_path} is {file_type}")
        # Add more processing logic here

# Example usage
if __name__ == "__main__":
    processor = DocumentProcessor()
    print(processor.extract_text("example.pdf"))